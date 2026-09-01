"""PDF/DOCX export. Every export is single-column, plain-text-styled, with

no tables, text boxes, or images - the layout choices that most commonly
break real ATS parsers - so a CV export is ATS-compatible by
construction, not just by claim. Narrative documents (SOP, personal
statement, motivation letter, study plan, research proposal, fellowship
essays) render their ``content["body"]`` text as simple paragraphs.
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.models.premium_documents import DocumentKind

_CV_KINDS = frozenset({DocumentKind.cv_academic, DocumentKind.cv_professional, DocumentKind.cv_scholarship})


def _cv_sections(content: dict) -> list[tuple[str, list[dict]]]:
    return [
        ("Education", content.get("education", []) or []),
        ("Experience", content.get("experience", []) or []),
        ("Projects", content.get("projects", []) or []),
        ("Publications", content.get("publications", []) or []),
        ("Awards", content.get("awards", []) or []),
        ("Leadership & Community", content.get("leadership", []) or []),
    ]


def _entry_text(entry: dict) -> str:
    name = entry.get("degree") or entry.get("role") or entry.get("title") or ""
    org = entry.get("institution") or entry.get("organization") or ""
    span = " - ".join(part for part in (entry.get("start_date", ""), entry.get("end_date", "")) if part)
    header = ", ".join(part for part in (name, org) if part)
    if span:
        header = f"{header} ({span})"
    description = entry.get("description") or ""
    return f"{header}: {description}" if description else header


def export_cv_pdf(title: str, content: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
        topMargin=0.75 * inch, bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle("CVHeading", parent=styles["Heading2"], spaceBefore=10, spaceAfter=4)
    story = [Paragraph(content.get("full_name") or title, styles["Title"])]
    contact = content.get("contact") or {}
    contact_line = " | ".join(v for v in contact.values() if v)
    if contact_line:
        story.append(Paragraph(contact_line, styles["Normal"]))
    if content.get("summary"):
        story.append(Spacer(1, 8))
        story.append(Paragraph("Summary", heading_style))
        story.append(Paragraph(content["summary"], styles["Normal"]))
    for label, entries in _cv_sections(content):
        if not entries:
            continue
        story.append(Paragraph(label, heading_style))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(_entry_text(entry), styles["Normal"])) for entry in entries],
                bulletType="bullet",
            )
        )
    skills = content.get("skills") or []
    if skills:
        story.append(Paragraph("Skills", heading_style))
        story.append(Paragraph(", ".join(skills), styles["Normal"]))
    doc.build(story)
    return buffer.getvalue()


def export_cv_docx(title: str, content: dict) -> bytes:
    document = DocxDocument()
    document.add_heading(content.get("full_name") or title, level=0)
    contact = content.get("contact") or {}
    contact_line = " | ".join(v for v in contact.values() if v)
    if contact_line:
        document.add_paragraph(contact_line)
    if content.get("summary"):
        document.add_heading("Summary", level=1)
        document.add_paragraph(content["summary"])
    for label, entries in _cv_sections(content):
        if not entries:
            continue
        document.add_heading(label, level=1)
        for entry in entries:
            document.add_paragraph(_entry_text(entry), style="List Bullet")
    skills = content.get("skills") or []
    if skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(", ".join(skills))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_narrative_pdf(title: str, content: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER, leftMargin=1 * inch, rightMargin=1 * inch,
        topMargin=1 * inch, bottomMargin=1 * inch,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    body = content.get("body", "")
    for paragraph in body.split("\n\n"):
        cleaned = paragraph.strip()
        if cleaned:
            story.append(Paragraph(cleaned.replace("\n", "<br/>"), styles["Normal"]))
            story.append(Spacer(1, 8))
    doc.build(story)
    return buffer.getvalue()


def export_narrative_docx(title: str, content: dict) -> bytes:
    document = DocxDocument()
    document.add_heading(title, level=0)
    body = content.get("body", "")
    for paragraph in body.split("\n\n"):
        cleaned = paragraph.strip()
        if cleaned:
            run_paragraph = document.add_paragraph()
            run = run_paragraph.add_run(cleaned)
            run.font.size = Pt(11)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_document(kind: DocumentKind, title: str, content: dict, fmt: str) -> tuple[bytes, str]:
    """Returns (bytes, content_type). ``fmt`` is "pdf" or "docx"."""
    is_cv = kind in _CV_KINDS
    if fmt == "pdf":
        data = export_cv_pdf(title, content) if is_cv else export_narrative_pdf(title, content)
        return data, "application/pdf"
    if fmt == "docx":
        data = export_cv_docx(title, content) if is_cv else export_narrative_docx(title, content)
        return data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise ValueError(f"Unsupported export format: {fmt}")
